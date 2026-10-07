"""
ATS Resume Checker
------------------
Upload a resume (PDF / DOCX / TXT), optionally paste a job description, and get:
  * an estimated ATS score (0-100) with a category breakdown
  * strengths and weaknesses
  * missing keywords
  * prioritised improvements and rewritten bullet points

UI: Streamlit    AI: Google Gemini Flash (via the google-genai SDK)
"""

import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
MODEL_OPTIONS = [
    "gemini-flash-latest",
    "gemini-3.5-flash-lite",
    "gemini-3.1-flash-lite",
    "gemini-3.5-flash",
    "gemini-3.6-flash",
]
MAX_RESUME_CHARS = 15000   # keeps prompts small and fast
MAX_JD_CHARS = 6000
MIN_TEXT_CHARS = 200       # below this we assume a scanned / image-only file

# Weights for the overall score (must sum to 1.0).
CATEGORY_WEIGHTS = {
    "keywords_and_skills": 0.30,
    "experience_and_impact": 0.25,
    "formatting_and_structure": 0.20,
    "readability_and_grammar": 0.10,
    "contact_information": 0.10,
    "education_and_certifications": 0.05,
}

CATEGORY_LABELS = {
    "keywords_and_skills": "Keywords & Skills",
    "experience_and_impact": "Experience & Impact",
    "formatting_and_structure": "Formatting & Structure",
    "readability_and_grammar": "Readability & Grammar",
    "contact_information": "Contact Information",
    "education_and_certifications": "Education & Certifications",
}


# --------------------------------------------------------------------------- #
# File reading
# --------------------------------------------------------------------------- #
def extract_text(uploaded_file) -> str:
    """Return plain text from an uploaded PDF, DOCX or TXT file."""
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected. Please upload an unlocked copy.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(uploaded_file)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Many resumes keep content inside tables (two-column layouts)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return uploaded_file.read().decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# --------------------------------------------------------------------------- #
# Prompt + model call
# --------------------------------------------------------------------------- #
def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n\"\"\"\n{job_description[:MAX_JD_CHARS]}\n\"\"\"\n"
        if job_description.strip()
        else "No job description was provided. Judge the resume against general ATS best practices "
             "and common expectations for the role it appears to target.\n"
    )

    return f"""You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.
Evaluate the resume below as an ATS would, then as a recruiter would.

Scoring rules:
- Score each category from 0 to 100. Be strict and realistic; most resumes score between 40 and 85.
- Only judge what is visible in the extracted text. Do not invent facts about the candidate.
- If a job description is given, "keywords_and_skills" must reflect how well the resume matches it,
  and "missing_keywords" must list important terms from the job description that the resume lacks.
- If no job description is given, "missing_keywords" should list commonly expected skills/keywords
  for the candidate's apparent target role.
- "improvements" must be specific and actionable, ordered from highest to lowest priority.
- "bullet_rewrites" must take REAL bullet points from the resume and improve them with strong action
  verbs and measurable impact. Do not invent numbers: use placeholders like [X%] where a metric is needed.

Return ONLY valid JSON (no markdown, no commentary) with exactly this structure:
{{
  "category_scores": {{
    "keywords_and_skills": <int 0-100>,
    "experience_and_impact": <int 0-100>,
    "formatting_and_structure": <int 0-100>,
    "readability_and_grammar": <int 0-100>,
    "contact_information": <int 0-100>,
    "education_and_certifications": <int 0-100>
  }},
  "jd_match_score": <int 0-100, or null if no job description was provided>,
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<string>", "..."],
  "weaknesses": ["<string>", "..."],
  "missing_keywords": ["<string>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "<resume section>", "issue": "<what is wrong>", "suggestion": "<how to fix it>"}}
  ],
  "bullet_rewrites": [
    {{"original": "<bullet from resume>", "improved": "<rewritten bullet>"}}
  ]
}}

{jd_block}
RESUME:
\"\"\"
{resume_text[:MAX_RESUME_CHARS]}
\"\"\"
"""


def parse_model_json(raw: str) -> dict:
    """Parse JSON from a model reply, tolerating code fences and stray text."""
    if not raw or not raw.strip():
        raise ValueError("The model returned an empty response.")

    text = raw.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end > start:
            try:
                return json.loads(text[start : end + 1])
            except json.JSONDecodeError:
                pass
    raise ValueError("Could not read the model's response as JSON. Please try again.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalize_result(data: dict) -> dict:
    """Validate the model output and compute the overall score deterministically."""
    raw_scores = data.get("category_scores") or {}
    scores = {key: _clamp(raw_scores.get(key)) for key in CATEGORY_WEIGHTS}

    overall = round(sum(scores[k] * w for k, w in CATEGORY_WEIGHTS.items()))

    jd_match = data.get("jd_match_score")
    jd_match = None if jd_match is None else _clamp(jd_match)

    improvements = []
    for item in data.get("improvements") or []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "Medium")).strip().capitalize()
        if priority not in ("High", "Medium", "Low"):
            priority = "Medium"
        improvements.append(
            {
                "priority": priority,
                "section": str(item.get("section", "General")).strip() or "General",
                "issue": str(item.get("issue", "")).strip(),
                "suggestion": str(item.get("suggestion", "")).strip(),
            }
        )
    order = {"High": 0, "Medium": 1, "Low": 2}
    improvements.sort(key=lambda i: order[i["priority"]])

    rewrites = []
    for item in data.get("bullet_rewrites") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append(
                {"original": str(item["original"]).strip(), "improved": str(item["improved"]).strip()}
            )

    return {
        "overall_score": overall,
        "category_scores": scores,
        "jd_match_score": jd_match,
        "summary": str(data.get("summary", "")).strip(),
        "strengths": _str_list(data.get("strengths")),
        "weaknesses": _str_list(data.get("weaknesses")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "improvements": improvements,
        "bullet_rewrites": rewrites,
    }


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str = "") -> dict:
    """Send the resume to Gemini and return a validated result dict."""
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,  # low temperature = more consistent scores
            response_mime_type="application/json",
        ),
    )
    return normalize_result(parse_model_json(response.text))


# --------------------------------------------------------------------------- #
# UI helpers
# --------------------------------------------------------------------------- #
def score_color(score: int) -> str:
    if score >= 80:
        return "🟢"
    if score >= 60:
        return "🟡"
    return "🔴"


def score_verdict(score: int) -> str:
    if score >= 80:
        return "Strong - likely to pass most ATS filters."
    if score >= 60:
        return "Decent - some improvements will boost your chances."
    return "Needs work - significant improvements recommended."


def get_api_key(sidebar_value: str) -> str:
    """Priority: Streamlit secrets -> environment variable -> sidebar input."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass  # no secrets file locally
    return os.getenv("GEMINI_API_KEY") or sidebar_value


def friendly_error(exc: Exception) -> str:
    msg = str(exc)
    low = msg.lower()
    if "api key" in low or "api_key" in low or "permission_denied" in low or "unauthenticated" in low:
        return "Your Gemini API key looks invalid or lacks permission. Please check it."
    if "429" in low or "quota" in low or "rate limit" in low or "resource_exhausted" in low:
        return "Rate limit or quota reached on the Gemini API. Wait a minute and try again."
    if "404" in low or "is not found" in low or "not_found" in low:
        return "That model name was not found. Try a different model in the sidebar."
    return f"Something went wrong: {msg}"


# --------------------------------------------------------------------------- #
# Streamlit app
# --------------------------------------------------------------------------- #
def main():
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")

    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume and get an estimated ATS score with concrete ways to improve it.")

    with st.sidebar:
        st.header("⚙️ Settings")
        sidebar_key = st.text_input(
            "Gemini API key",
            type="password",
            help="Get a free key at https://aistudio.google.com/apikey. "
                 "Not needed if the app owner configured one.",
        )
        model = st.selectbox("Gemini model", MODEL_OPTIONS, index=0)
        st.markdown("---")
        st.info(
            "The score is an **AI-based estimate**, not the output of a real ATS. "
            "Use it as guidance for improving your resume."
        )

    api_key = get_api_key(sidebar_key)

    col_left, col_right = st.columns(2)
    with col_left:
        uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    with col_right:
        job_description = st.text_area(
            "Job description (optional, but gives a better keyword match)",
            height=160,
            placeholder="Paste the job description here...",
        )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            st.stop()

        try:
            resume_text = extract_text(uploaded)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            st.stop()

        if len(resume_text) < MIN_TEXT_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned image, "
                "ATS systems cannot read it either - export a text-based PDF or DOCX instead."
            )
            st.stop()

        with st.spinner("Analyzing your resume with Gemini..."):
            try:
                result = analyze_resume(api_key, model, resume_text, job_description)
            except Exception as exc:
                st.error(friendly_error(exc))
                st.stop()

        render_results(result, has_jd=bool(job_description.strip()))
        with st.expander("View extracted resume text"):
            st.text(resume_text)


def render_results(result: dict, has_jd: bool):
    overall = result["overall_score"]

    st.divider()
    top_left, top_right = st.columns([1, 2])
    with top_left:
        st.metric("Estimated ATS score", f"{overall} / 100")
        st.progress(overall / 100)
        st.write(f"{score_color(overall)} {score_verdict(overall)}")
        if has_jd and result["jd_match_score"] is not None:
            st.metric("Job description match", f"{result['jd_match_score']} / 100")
    with top_right:
        st.subheader("Summary")
        st.write(result["summary"] or "No summary returned.")

    st.subheader("Score breakdown")
    cols = st.columns(3)
    for i, (key, label) in enumerate(CATEGORY_LABELS.items()):
        score = result["category_scores"][key]
        with cols[i % 3]:
            st.write(f"**{label}** - {score_color(score)} {score}/100")
            st.progress(score / 100)

    s_col, w_col = st.columns(2)
    with s_col:
        st.subheader("✅ Strengths")
        for item in result["strengths"] or ["None identified."]:
            st.markdown(f"- {item}")
    with w_col:
        st.subheader("⚠️ Weaknesses")
        for item in result["weaknesses"] or ["None identified."]:
            st.markdown(f"- {item}")

    st.subheader("🔑 Missing keywords")
    if result["missing_keywords"]:
        st.write(" ".join(f"`{kw}`" for kw in result["missing_keywords"]))
    else:
        st.write("No major keyword gaps found.")

    st.subheader("🛠️ Recommended improvements")
    icons = {"High": "🔴", "Medium": "🟡", "Low": "🟢"}
    if result["improvements"]:
        for imp in result["improvements"]:
            with st.expander(f"{icons[imp['priority']]} {imp['priority']} - {imp['section']}"):
                st.markdown(f"**Issue:** {imp['issue']}")
                st.markdown(f"**Fix:** {imp['suggestion']}")
    else:
        st.write("No improvements suggested.")

    if result["bullet_rewrites"]:
        st.subheader("✍️ Bullet point rewrites")
        for rw in result["bullet_rewrites"]:
            st.markdown(f"**Before:** {rw['original']}")
            st.markdown(f"**After:** {rw['improved']}")
            st.markdown("---")


if __name__ == "__main__":
    main()
